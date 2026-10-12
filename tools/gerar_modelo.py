"""Gera os modelos de importação em modelos/: uma planilha .xlsx com cabeçalhos coloridos
(vermelho = obrigatório, amarelo = opcional importante, sem cor = opcional) e um CSV
só com cabeçalho para cada tabela.

Uso: python tools/gerar_modelo.py   (requer openpyxl)
Mantenha TEMPLATE alinhado com a constante TEMPLATE em index.html.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REQ, IMP, OPT = "obrigatório", "importante", "opcional"

TEMPLATE = {
    "customer": [
        ("id", REQ, "ID da linha de cliente. Igual a kpi.id_customer."),
        ("dt_register", REQ, "Data de início do cliente/contrato. Define a coorte."),
        ("dt_cancel", REQ, "Data do cancelamento. Vazio = cliente ativo."),
        ("name", IMP, "Nome exibido nas tabelas. Ex.: Ambev - Produto 1."),
        ("hierarchy_type", IMP, "O que a linha representa: grupo, organizacao ou produto. Linhas de grupo ou organização com linhas abaixo só dão nome à hierarquia."),
        ("id_group", IMP, "Hierarquia: ID do grupo econômico."),
        ("group", IMP, "Hierarquia: nome do grupo econômico."),
        ("id_org", IMP, "Hierarquia: ID da organização (o próprio cliente)."),
        ("id_product", IMP, "Hierarquia: ID do produto contratado, quando o cliente tem uma linha por produto."),
        ("product", IMP, "Hierarquia: nome do produto."),
        ("industry", IMP, "Perfil: setor."),
        ("size", IMP, "Perfil: porte."),
        ("state", IMP, "Perfil: UF."),
        ("stage", IMP, "Perfil: etapa do ciclo de vida."),
        ("cs_owner", IMP, "Perfil: nome do CSM responsável (texto, como os demais campos de perfil)."),
        ("cancel_tag", IMP, "Motivo do cancelamento (tag)."),
        ("cancel_description", IMP, "Descrição do motivo."),
        ("mrr", IMP, "Receita mensal da linha. Ex.: custom_fields->>'mrr'."),
        ("mrr_12m_atras", OPT, "MRR da linha 12 meses antes. Com ele o relatório calcula o NRR."),
        ("cnpj", OPT, "CNPJ da organização."),
        ("id_legacy", OPT, "ID no sistema anterior."),
        ("name_contract", OPT, "Nome do contrato."),
        ("city", OPT, "Cidade."),
        ("country", OPT, "País."),
        ("support_owner", OPT, "Responsável de suporte. Vira perfil se tiver poucos valores."),
        ("sales_owner", OPT, "Vendedor. Vira perfil se tiver poucos valores."),
        ("onboarding_owner", OPT, "Responsável pelo onboarding. Vira perfil se tiver poucos valores."),
        ("dt_stage", OPT, "Data da etapa atual."),
    ],
    "kpi": [
        ("id_customer", REQ, "ID da linha de cliente (customer.id)."),
        ("type", REQ, "ID do tipo de KPI (kpi_type.id)."),
        ("value", REQ, "Valor do KPI."),
        ("ref_date", IMP, "Data da leitura. Sem ela, vale a média das linhas do cliente."),
        ("kpi_name", IMP, "Nome do KPI (kpi_type.name). Sem ele, aparece como KPI 1, KPI 2…"),
        ("inv_proportional", IMP, "kpi_type.inv_proportional. true = valor maior é pior."),
        ("unit", OPT, "Unidade (kpi_type.unit)."),
    ],
    "cancelamentos": [
        ("id_customer", REQ, "ID da linha de cliente. Só use se o cancelamento não estiver em customer."),
        ("dt_cancel", REQ, "Data do cancelamento."),
        ("cancel_tag", IMP, "Motivo do cancelamento."),
        ("mrr_perdido", OPT, "MRR perdido, se diferente do mrr em customer."),
    ],
}

FILL = {REQ: PatternFill("solid", fgColor="F4B6B6"), IMP: PatternFill("solid", fgColor="FBE29A"), OPT: None}
THIN = Side(style="thin", color="BFC7C3")


def main():
    out = Path(__file__).resolve().parent.parent / "modelos"
    out.mkdir(exist_ok=True)
    wb = Workbook()
    leia = wb.active
    leia.title = "Leia-me"
    leia.append(["Modelo de importação · Análise de Churn por Coorte"])
    leia["A1"].font = Font(bold=True, size=14)
    leia.append([])
    leia.append(["Cor do cabeçalho", "Significado"])
    for cell in leia[3]:
        cell.font = Font(bold=True)
    for kind, text in [
        (REQ, "Obrigatória para a análise."),
        (IMP, "Opcional, mas importante: usada em filtros, hierarquia e comparações."),
        (OPT, "Totalmente opcional."),
    ]:
        leia.append([kind, text])
        if FILL[kind]:
            leia.cell(leia.max_row, 1).fill = FILL[kind]
    leia.append([])
    leia.append(["Exporte cada aba como CSV (com cabeçalho) e anexe no cartão de mesmo nome no relatório."])
    leia.append(["A aba cancelamentos só é necessária se o cancelamento não estiver em customer."])
    leia.append([])
    leia.append(["Tabela", "Coluna", "Tipo", "Descrição"])
    for cell in leia[leia.max_row]:
        cell.font = Font(bold=True)
    for table, cols in TEMPLATE.items():
        for name, kind, desc in cols:
            leia.append([table, name, kind, desc])
            if FILL[kind]:
                leia.cell(leia.max_row, 3).fill = FILL[kind]
    for col, width in zip("ABCD", (18, 22, 14, 80)):
        leia.column_dimensions[col].width = width

    for table, cols in TEMPLATE.items():
        ws = wb.create_sheet(table)
        ws.append([c[0] for c in cols])
        for i, (name, kind, desc) in enumerate(cols, start=1):
            cell = ws.cell(1, i)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(bottom=THIN, right=THIN)
            if FILL[kind]:
                cell.fill = FILL[kind]
            cell.comment = Comment(f"{kind}: {desc}", "modelo")
            ws.column_dimensions[cell.column_letter].width = max(14, len(name) + 4)
        ws.freeze_panes = "A2"
        with open(out / f"{table}.csv", "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow([c[0] for c in cols])

    wb.save(out / "modelo_churn.xlsx")
    print("modelos gerados em", out)


if __name__ == "__main__":
    main()
